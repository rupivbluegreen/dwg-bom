"""Stage 6 - write the results: <name>_BOM.xlsx for people, <name>_bom.json for other tools."""
from __future__ import annotations

import datetime as dt
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT = "Arial"
F_BASE = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=FONT, size=14, bold=True)
F_GREY = Font(name=FONT, size=10, color="808080")
F_BLUE = Font(name=FONT, size=10, color="0000FF")
FILL_HEAD = PatternFill("solid", start_color="305496")
FILL_INPUT = PatternFill("solid", start_color="FFFF00")
FILL_ERR = PatternFill("solid", start_color="F8CBAD")
THIN = Border(bottom=Side(style="thin", color="BFBFBF"))
WRAP = Alignment(wrap_text=True, vertical="top")

BOM_COLUMNS = [  # header, key, width
    ("Item", "item", 6), ("Category", "category", 20), ("Description", "description", 38),
    ("Designation", "designation", 24), ("kg/m", "mass_kg_m", 8), ("Unit", "unit", 6),
    ("Qty", None, 7), ("Length each (m)", None, 10), ("Total length (m)", None, 10),
    ("Total mass (kg)", None, 11), ("Callouts on sheet", "callouts", 9),
    ("Drawing evidence", "evidence", 30), ("Remarks", "remarks", 40), ("Status", "status", 26),
    ("Scope", None, 10), ("Source ids", "sources", 22),
]


def write_outputs(result, extraction: dict, out_dir: Path, stem: str, warnings: list[str]) -> tuple[Path, Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    xlsx = out_dir / f"{stem}_BOM.xlsx"
    js = out_dir / f"{stem}_bom.json"

    wb = Workbook()
    _bom_sheet(wb.active, result, extraction)
    _table_sheet(wb.create_sheet("Flags"), ["Severity", "Type", "Source id", "Message", "Note text"],
                 [[f["severity"], f["type"], f["source_id"], f["message"], f["text"]] for f in result.flags],
                 [9, 22, 10, 55, 60], highlight_col=0, highlight_value="error")
    _table_sheet(wb.create_sheet("Notes"), ["Id", "Layer", "Space", "Block", "Text", "Classified as"],
                 [[n["id"], n["layer"], n["space"], n["block"] or "", n["text"], n["classified_as"]] for n in result.notes],
                 [10, 18, 10, 14, 60, 50])
    _info_sheet(wb.create_sheet("Drawing info"), extraction, warnings)
    wb.save(xlsx)

    payload = {"source": extraction.get("source"), "title_block": extraction.get("title_block"),
               "generated": dt.datetime.now().isoformat(timespec="seconds"),
               "bom": result.rows, "flags": result.flags, "notes": result.notes}
    js.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
    return xlsx, js


def _bom_sheet(ws, result, extraction):
    ws.title = "BOM"
    tb = _flatten_title_block(extraction.get("title_block", {}))
    src = extraction.get("source", {})
    ws["A1"] = "Bill of materials - DRAFT (generated from drawing, check before use)"
    ws["A1"].font = F_TITLE
    meta = [
        ("Project", " / ".join(v for k, v in tb.items() if k.startswith("PROJECT_NAME") or k.startswith("PROJECT_PHASE"))),
        ("Sheet", " ".join(v for k, v in tb.items() if "SHEETNUMBER" in k or "SHEET_CONTENT" in k)),
        ("Source file", src.get("file", "")),
        ("Generated", dt.datetime.now().strftime("%Y-%m-%d %H:%M")),
    ]
    for i, (k, v) in enumerate(meta, start=2):
        ws.cell(i, 1, k).font = F_BOLD
        ws.cell(i, 3, v).font = F_BASE
    ws["A7"] = ("Yellow cells are for you to fill in (Qty, Length each, missing kg/m). Totals update automatically. "
                "Example: Qty 2 and Length each 3.25 gives Total length 6.5 m; with 39 kg/m, Total mass 253.5 kg. "
                "'Callouts on sheet' counts notes pointing at an item - it is NOT a quantity.")
    ws["A7"].font = F_GREY
    ws.merge_cells("A7:P7")
    ws["A7"].alignment = WRAP
    ws.row_dimensions[7].height = 30

    hr = 9
    for c, (head, _, width) in enumerate(BOM_COLUMNS, start=1):
        cell = ws.cell(hr, c, head)
        cell.font, cell.fill, cell.alignment = F_HEAD, FILL_HEAD, Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.row_dimensions[hr].height = 30

    r = hr
    for row in result.rows:
        r += 1
        for c, (head, key, _) in enumerate(BOM_COLUMNS, start=1):
            if key:
                val = row.get(key)
                ws.cell(r, c, "" if val is None else val)
        ws.cell(r, 15, "Reference" if row["reference_only"] else "BOM")
        # formulas: I = G*H, J = E*I
        ws.cell(r, 9, f'=IF(AND(ISNUMBER(G{r}),ISNUMBER(H{r})),G{r}*H{r},"")')
        ws.cell(r, 10, f'=IF(AND(ISNUMBER(E{r}),ISNUMBER(I{r})),E{r}*I{r},"")')
        for c in range(1, len(BOM_COLUMNS) + 1):
            cell = ws.cell(r, c)
            cell.font = F_GREY if row["reference_only"] else F_BASE
            cell.alignment, cell.border = WRAP, THIN
        for c in (7, 8):
            ws.cell(r, c).fill = FILL_INPUT
            ws.cell(r, c).font = F_BLUE
        if row["mass_kg_m"] is None and row["category"] == "Structural steel" and row["unit"] == "m":
            ws.cell(r, 5).fill = FILL_INPUT
        elif row["mass_kg_m"] is not None:
            ws.cell(r, 5).font = F_BLUE
        ws.cell(r, 5).number_format = "0.0"
        ws.cell(r, 8).number_format = "0.000"
        ws.cell(r, 9).number_format = "0.000"
        ws.cell(r, 10).number_format = "#,##0.0"

    total_row = r + 2
    ws.cell(total_row, 9, "Total steel mass (kg)").font = F_BOLD
    ws.cell(total_row, 10, f"=SUM(J{hr + 1}:J{max(r, hr + 1)})").font = F_BOLD
    ws.cell(total_row, 10).number_format = "#,##0.0"
    ws.freeze_panes = ws.cell(hr + 1, 4)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.print_title_rows = f"{hr}:{hr}"
    ws.auto_filter.ref = f"A{hr}:{get_column_letter(len(BOM_COLUMNS))}{max(r, hr + 1)}"


def _table_sheet(ws, headers, rows, widths, highlight_col=None, highlight_value=None):
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(1, c, h)
        cell.font, cell.fill = F_HEAD, FILL_HEAD
        ws.column_dimensions[get_column_letter(c)].width = widths[c - 1]
    for r, row in enumerate(rows, start=2):
        for c, v in enumerate(row, start=1):
            cell = ws.cell(r, c, v)
            cell.font, cell.alignment, cell.border = F_BASE, WRAP, THIN
        if highlight_col is not None and row[highlight_col] == highlight_value:
            for c in range(1, len(headers) + 1):
                ws.cell(r, c).fill = FILL_ERR
    if not rows:
        ws.cell(2, 1, "(none)").font = F_GREY
    ws.freeze_panes = "A2"


def _info_sheet(ws, ex, warnings):
    ws.column_dimensions["A"].width = 28
    for col in "BCDE":
        ws.column_dimensions[col].width = 26
    r = 1

    def section(title):
        nonlocal r
        r += 1 if r > 1 else 0
        ws.cell(r, 1, title).font = F_TITLE
        r += 1

    def line(*vals, bold=False):
        nonlocal r
        for c, v in enumerate(vals, start=1):
            ws.cell(r, c, v).font = F_BOLD if bold else F_BASE
        r += 1

    section("Source")
    for k, v in (ex.get("source") or {}).items():
        if k != "audit_issues":
            line(k, "" if v is None else str(v))
    section("Title block")
    for k, v in _flatten_title_block(ex.get("title_block", {})).items():
        line(k, v)
    section("Viewports")
    line("Layout", "Label", "Scale", "Scale factor", bold=True)
    for v in ex.get("viewports", []):
        line(v.get("layout"), v.get("label", ""), v.get("scale"), v.get("scale_factor"))
    section("Blocks (inserts, nested blocks followed)")
    line("Block", "Count", "Layers", "Spaces", bold=True)
    for b in ex.get("blocks", []):
        line(b["name"], "unknown" if b.get("count") is None else b["count"], ", ".join(b.get("layers", [])),
             ", ".join(b.get("spaces", [])))
    section("Linework on structural layers (drawn lines, not member lengths)")
    line("Layer", "Entities", "Length (drawing units)", "Length (m)", bold=True)
    for l in ex.get("linework", []):
        line(l["layer"], l["entities"], l["length_units"], l["length_m"])
    section("Conversion warnings / audit issues")
    for w in list(warnings) + list((ex.get("source") or {}).get("audit_issues", [])):
        line(w)


def _flatten_title_block(tb: dict) -> dict:
    out = {}
    for block, attrs in (tb or {}).items():
        for tag, val in attrs.items():
            if val:
                out[tag if tag not in out else f"{block}.{tag}"] = val
    return out
